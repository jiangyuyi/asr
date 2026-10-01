"""翻译相关的导出测试。

重点是**对齐**：译文列表与片段列表等长且同序，任何一处 0/1 基混用都会让整列
译文错位一行，肉眼扫一遍表格很容易漏掉，所以这里逐行断言。
"""
from __future__ import annotations

import json

import pytest

from asr_mm import export
from asr_mm.srt import Cue

SEGMENTS = [
    Cue(1.0, 2.0, "第一句中文。"),
    Cue(2.0, 3.5, "第二句中文。"),
    Cue(3.5, 5.0, "第三句中文。"),
]
EN = ["first english line.", "second english line.", "third english line."]
JA = ["一つ目の日本語。", "二つ目の日本語。", "三つ目の日本語。"]


def test_xlsx_translation_columns_align_with_rows(tmp_path):
    from openpyxl import load_workbook
    p = tmp_path / "t.xlsx"
    export.write(p, SEGMENTS, "xlsx", lang="zh",
                 translations={"en": EN, "ja": JA})
    ws = load_workbook(p).active
    rows = list(ws.iter_rows(values_only=True))

    assert rows[0] == ("序号", "开始", "结束", "开始(秒)", "时长(秒)",
                       "内容", "英文", "日文")
    # 序号从 1 开始，译文却必须跟着同序的第 1/2/3 行内容走
    for n, row in enumerate(rows[1:]):
        assert row[0] == n + 1
        assert row[5] == SEGMENTS[n].text
        assert row[6] == EN[n], f"第 {n+1} 行英文错位: {row[6]!r} != {EN[n]!r}"
        assert row[7] == JA[n], f"第 {n+1} 行日文错位: {row[7]!r} != {JA[n]!r}"


def test_xlsx_without_translations_keeps_six_columns(tmp_path):
    from openpyxl import load_workbook
    p = tmp_path / "plain.xlsx"
    export.write(p, SEGMENTS, "xlsx", lang="zh")
    rows = list(load_workbook(p).active.iter_rows(values_only=True))
    assert len(rows[0]) == 6
    # 空译文不该凭空多出两列
    p2 = tmp_path / "blank.xlsx"
    export.write(p2, SEGMENTS, "xlsx", lang="zh",
                 translations={"en": ["", "", ""], "ja": ["", "", ""]})
    assert len(list(load_workbook(p2).active.iter_rows(values_only=True))[0]) == 6


def test_xlsx_translation_headers_follow_interface_language(tmp_path):
    from openpyxl import load_workbook
    p = tmp_path / "en.xlsx"
    export.write(p, SEGMENTS, "xlsx", lang="en", translations={"en": EN, "ja": JA})
    header = list(load_workbook(p).active.iter_rows(values_only=True))[0]
    assert header[5] == "Text"
    assert header[6] == "English"
    assert header[7] == "Japanese"

    # 只译了日文时，那一列直接接在原文后面，而不是留一个空的英文列
    p2 = tmp_path / "en_ja_only.xlsx"
    export.write(p2, SEGMENTS, "xlsx", lang="en", translations={"ja": JA})
    header2 = list(load_workbook(p2).active.iter_rows(values_only=True))[0]
    assert header2 == ("#", "Start", "End", "Start (s)", "Duration (s)",
                       "Text", "Japanese")

    p3 = tmp_path / "ja_ui.xlsx"
    export.write(p3, SEGMENTS, "xlsx", lang="ja", translations={"en": EN, "ja": JA})
    header3 = list(load_workbook(p3).active.iter_rows(values_only=True))[0]
    assert header3[6] == "英訳"
    assert header3[7] == "日訳"


def test_srt_puts_every_language_in_the_same_cue(tmp_path):
    p = tmp_path / "t.srt"
    export.write(p, SEGMENTS, "srt", translations={"en": EN, "ja": JA})
    text = p.read_text(encoding="utf-8")
    blocks = [b for b in text.strip().split("\n\n")]
    assert len(blocks) == 3
    for n, block in enumerate(blocks):
        lines = block.splitlines()
        assert lines[0] == str(n + 1)
        assert lines[2] == SEGMENTS[n].text
        assert lines[3] == EN[n]
        assert lines[4] == JA[n]


def test_txt_puts_every_language_on_one_line(tmp_path):
    p = tmp_path / "t.txt"
    export.write(p, SEGMENTS, "txt", translations={"en": EN, "ja": JA})
    lines = p.read_text(encoding="utf-8").splitlines()
    assert len(lines) == 3
    for n, line in enumerate(lines):
        assert SEGMENTS[n].text in line
        assert EN[n] in line
        assert JA[n] in line


def test_json_gains_a_translations_block(tmp_path):
    p = tmp_path / "t.json"
    payload = {"segments": [{"text": s.text} for s in SEGMENTS]}
    export.write(p, SEGMENTS, "json", json_payload=payload,
                 translations={"en": EN, "ja": JA})
    data = json.loads(p.read_text(encoding="utf-8"))
    assert data["translations"] == {"en": EN, "ja": JA}
    # 原有字段不能被挤掉
    assert "segments" in data


def test_json_without_translations_is_unchanged(tmp_path):
    p = tmp_path / "t.json"
    payload = {"segments": [{"text": s.text} for s in SEGMENTS]}
    export.write(p, SEGMENTS, "json", json_payload=payload)
    assert "translations" not in json.loads(p.read_text(encoding="utf-8"))


def test_shorter_translation_list_is_padded_not_shifted():
    """译文条数少于片段数时只允许缺尾部，不允许整体错位。"""
    from openpyxl import load_workbook
    import tempfile
    from pathlib import Path
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "short.xlsx"
        export.write(p, SEGMENTS, "xlsx", translations={"en": EN[:2]})
        rows = list(load_workbook(p).active.iter_rows(values_only=True))
        assert rows[1][6] == EN[0]
        assert rows[2][6] == EN[1]
        assert rows[3][6] is None


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))
