"""Unit tests for output formats and the content-language model check."""
import os
import sys

os.environ.setdefault("ASR_MM_HOME", os.path.join(os.path.dirname(__file__), ".testhome"))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest  # noqa: E402

from asr_mm import catalog, export  # noqa: E402
from asr_mm.srt import Cue  # noqa: E402
from asr_mm.transcribe import Segment  # noqa: E402

# 虚构的示例句。不要替换成真实课堂录像的转写内容——这些字符串会进公开仓库。
CUES = [
    Cue(15.3, 25.3, "请把这份材料复印三份，其中一部分留给我自己用。"),
    Cue(27.35, 37.35, "因为前面那条路正在维修，公交车今天临时改道绕行。"),
    Cue(37.35, 44.98, "你觉得这个价格合理吗？如果不合适可以再商量。"),
]
SEGMENTS = [Segment(c.start, c.end, c.text) for c in CUES]
PAYLOAD = {"source": "v.mp4", "segments": []}


# ------------------------------------------------------------------ dispatch

def test_all_four_formats_are_offered():
    assert set(export.FORMATS) == {"txt", "srt", "json", "xlsx"}


def test_every_format_has_an_extension():
    for fmt in export.FORMATS:
        assert fmt in export.EXTENSIONS


def test_guess_format_from_extension():
    assert export.guess_format("a.xlsx") == "xlsx"
    assert export.guess_format("a.SRT") == "srt"
    assert export.guess_format("a.txt") == "txt"
    assert export.guess_format("noextension") == "txt"
    assert export.guess_format("a.docx") == "txt"


def test_unknown_format_raises(tmp_path):
    with pytest.raises(export.ExportError):
        export.write(tmp_path / "a.xyz", SEGMENTS, "xyz")


# ------------------------------------------------------------------ txt/srt

def test_txt_has_one_line_per_cue(tmp_path):
    p = tmp_path / "a.txt"
    export.write(p, SEGMENTS, "txt")
    text = p.read_text(encoding="utf-8")
    for cue in CUES:
        assert cue.text in text


def test_srt_roundtrips(tmp_path):
    from asr_mm.srt import parse_srt
    p = tmp_path / "a.srt"
    export.write(p, SEGMENTS, "srt")
    back = parse_srt(p.read_text(encoding="utf-8"))
    assert len(back) == len(CUES)
    assert back[0].start == pytest.approx(15.3)


def test_json_keeps_the_payload(tmp_path):
    p = tmp_path / "a.json"
    export.write(p, SEGMENTS, "json", json_payload=PAYLOAD)
    import json
    assert json.loads(p.read_text(encoding="utf-8"))["source"] == "v.mp4"


# -------------------------------------------------------------------- xlsx

def test_xlsx_one_row_per_cue(tmp_path):
    from openpyxl import load_workbook
    p = tmp_path / "a.xlsx"
    export.write(p, SEGMENTS, "xlsx")
    ws = load_workbook(p).active
    assert ws.max_row == len(CUES) + 1          # header + one per cue
    assert ws.max_column == 6


def test_xlsx_columns_carry_timestamps_and_text(tmp_path):
    from openpyxl import load_workbook
    p = tmp_path / "a.xlsx"
    export.write(p, SEGMENTS, "xlsx", lang="zh")
    ws = load_workbook(p).active
    headers = [c.value for c in ws[1]]
    assert headers == ["序号", "开始", "结束", "开始(秒)", "时长(秒)", "内容"]
    row2 = [c.value for c in ws[2]]
    assert row2[0] == 1                        # index
    assert row2[1] == "00:00:15.300"           # start timecode
    assert row2[2] == "00:00:25.300"           # end timecode
    assert row2[3] == pytest.approx(15.3)       # numeric start
    assert row2[4] == pytest.approx(10.0)       # numeric duration
    assert row2[5] == CUES[0].text


def test_xlsx_headers_follow_the_interface_language(tmp_path):
    from openpyxl import load_workbook
    for lang, first in (("zh", "序号"), ("en", "#"), ("ja", "番号")):
        p = tmp_path / f"{lang}.xlsx"
        export.write(p, SEGMENTS, "xlsx", lang=lang)
        assert load_workbook(p).active["A1"].value == first


def test_xlsx_freezes_the_header_and_filters(tmp_path):
    from openpyxl import load_workbook
    p = tmp_path / "a.xlsx"
    export.write(p, SEGMENTS, "xlsx")
    ws = load_workbook(p).active
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None


def test_xlsx_accepts_plain_cues_as_well_as_segments(tmp_path):
    from openpyxl import load_workbook
    p = tmp_path / "a.xlsx"
    export.write(p, CUES, "xlsx")
    assert load_workbook(p).active.max_row == len(CUES) + 1


def test_xlsx_creates_missing_parent_directories(tmp_path):
    p = tmp_path / "deep" / "nested" / "a.xlsx"
    export.write(p, SEGMENTS, "xlsx")
    assert p.exists()


# -------------------------------------------------------------- write_all

def test_write_all_writes_each_format(tmp_path):
    written = export.write_all(tmp_path, "clip", SEGMENTS,
                               ["txt", "srt", "json", "xlsx"],
                               json_payload=PAYLOAD)
    assert [p.suffix for p in written] == [".txt", ".srt", ".json", ".xlsx"]
    assert all(p.exists() and p.stat().st_size > 0 for p in written)


def test_write_all_respects_the_stem(tmp_path):
    written = export.write_all(tmp_path, "我的视频_15s-45s", SEGMENTS, ["txt", "xlsx"])
    assert all(p.name.startswith("我的视频_15s-45s") for p in written)


# ------------------------------------------------------- language capability

def test_each_model_declares_the_languages_it_measured():
    nano = catalog.resolve_model("nano")
    para = catalog.resolve_model("paraformer")
    sense = catalog.resolve_model("sensevoice")
    assert set(nano.languages) == {"zh", "en", "ja"}
    assert para.languages == ("zh",)          # Japanese came back empty
    assert set(sense.languages) == {"zh", "en", "ja"}


def test_paraformer_warns_for_non_chinese():
    para = catalog.resolve_model("paraformer")
    assert para.warning_for("zh") == ""
    assert para.warning_for("ja") != ""
    assert para.warning_for("en") != ""


def test_auto_never_warns():
    for key in catalog.MODELS:
        spec = catalog.resolve_model(key)
        assert spec.warning_for("auto") == ""


def test_nano_never_warns():
    nano = catalog.resolve_model("nano")
    for code in ("zh", "en", "ja", "auto"):
        assert nano.warning_for(code) == ""


def test_sensevoice_marks_japanese_as_degraded():
    sense = catalog.resolve_model("sensevoice")
    assert "ja" in sense.degraded
    assert sense.warning_for("ja") == ""       # usable, just not clean


def test_supports_helper():
    assert catalog.resolve_model("paraformer").supports("zh")
    assert not catalog.resolve_model("paraformer").supports("ja")
    assert catalog.resolve_model("paraformer").supports("auto")


def test_warning_text_is_localised():
    from asr_mm.i18n import set_language
    para = catalog.resolve_model("paraformer")
    seen = set()
    for lang in ("zh", "en", "ja"):
        set_language(lang)
        seen.add(para.warning_for("ja"))
    set_language("zh")
    assert len(seen) == 3, "the warning must differ per language"


# ----------------------------------------------------------------- i18n keys

def test_new_keys_exist_in_all_languages():
    from asr_mm import i18n
    for lang in ("zh", "en", "ja"):
        missing = i18n.missing_keys(lang)
        assert not missing, f"{lang} missing: {missing}"


def test_format_help_mentions_xlsx_in_all_languages():
    from asr_mm import i18n
    for lang in ("zh", "en", "ja"):
        i18n.set_language(lang)
        assert "xlsx" in i18n.t("cli.format.help")
    i18n.set_language("zh")
