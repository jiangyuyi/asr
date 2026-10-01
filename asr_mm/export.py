"""Writing a transcript out.

Both the CLI and the GUI produced files themselves, and the logic had already
drifted (the GUI exported the table's current contents, the CLI exported the
raw segments). Everything lives here now so a format is implemented once and
both front ends agree on what it contains.

Excel is the format people actually paste into reports, so it gets a real
spreadsheet: one row per utterance, sortable numeric columns, a frozen header
and wrapped text. Not a renamed CSV.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable, Sequence

from .srt import Cue, format_ts, render_srt, render_txt

FORMATS = ("txt", "srt", "json", "xlsx")
EXTENSIONS = {"txt": ".txt", "srt": ".srt", "json": ".json", "xlsx": ".xlsx"}

# Excel column layout. Start/end are timecode strings because that is what a
# user needs to find the line in the video; the seconds columns are there for
# sorting and for anything that wants to compute with them.
XLSX_HEADERS_ZH = ["序号", "开始", "结束", "开始(秒)", "时长(秒)", "内容"]
XLSX_HEADERS_EN = ["#", "Start", "End", "Start (s)", "Duration (s)", "Text"]
XLSX_HEADERS_JA = ["番号", "開始", "終了", "開始(秒)", "長さ(秒)", "内容"]


class ExportError(RuntimeError):
    pass


def _headers(lang: str) -> list[str]:
    return {"en": XLSX_HEADERS_EN, "ja": XLSX_HEADERS_JA}.get(lang, XLSX_HEADERS_ZH)


def _as_cues(segments) -> list[Cue]:
    """Accept Segment objects or plain cues alike."""
    out = []
    for s in segments:
        if isinstance(s, Cue):
            out.append(s)
        else:
            out.append(Cue(s.start, s.end, s.text))
    return out


def write_txt(path: Path, segments, *, timestamps: bool = False,
              lang: str = "zh") -> None:
    path.write_text(render_txt(_as_cues(segments), timestamps), encoding="utf-8")


def write_srt(path: Path, segments) -> None:
    path.write_text(render_srt(_as_cues(segments)), encoding="utf-8")


def write_json(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                    encoding="utf-8")


def write_xlsx(path: Path, segments, *, lang: str = "zh") -> None:
    """One row per utterance.

    A missing openpyxl is a clear message rather than an ImportError traceback
    several frames deep.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ExportError(
            "导出 Excel 需要 openpyxl，请重新安装本程序。") from exc

    cues = _as_cues(segments)
    wb = Workbook()
    ws = wb.active
    ws.title = "Transcript"

    headers = _headers(lang)
    ws.append(headers)

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="374151")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, c in enumerate(cues, 1):
        ws.append([i, format_ts(c.start, comma=False), format_ts(c.end, comma=False),
                   round(c.start, 3), round(c.duration, 2), c.text])

    widths = [6, 14, 14, 11, 11, 70]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    text_col = len(headers)
    for row in range(2, len(cues) + 2):
        ws.cell(row=row, column=text_col).alignment = Alignment(
            wrap_text=True, vertical="top")
        for col in (1, 4, 5):
            ws.cell(row=row, column=col).alignment = Alignment(
                horizontal="right", vertical="top")

    ws.freeze_panes = "A2"
    if cues:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(cues) + 1}"
    ws.row_dimensions[1].height = 22

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write(path: Path, segments, fmt: str, *, timestamps: bool = False,
          lang: str = "zh", json_payload: dict | None = None) -> Path:
    """Dispatch to one writer. Returns the path written."""
    fmt = fmt.lower().lstrip(".")
    if fmt == "txt":
        write_txt(path, segments, timestamps=timestamps, lang=lang)
    elif fmt == "srt":
        write_srt(path, segments)
    elif fmt == "json":
        write_json(path, json_payload or {})
    elif fmt == "xlsx":
        write_xlsx(path, segments, lang=lang)
    else:
        raise ExportError(f"未知的导出格式: {fmt}")
    return path


def write_all(out_dir: Path, stem: str, segments, formats: Sequence[str], *,
              timestamps: bool = False, lang: str = "zh",
              json_payload: dict | None = None) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for fmt in formats:
        target = out_dir / f"{stem}{EXTENSIONS.get(fmt, '.' + fmt)}"
        written.append(write(target, segments, fmt, timestamps=timestamps,
                             lang=lang, json_payload=json_payload))
    return written


def export_dialog_filters(include_excel: bool = True) -> str:
    """QFileDialog filter string covering the formats we can write."""
    parts = ["文本文件 (*.txt)", "字幕文件 (*.srt)", "JSON (*.json)"]
    if include_excel:
        parts.insert(0, "Excel 工作簿 (*.xlsx)")
    parts.append("所有文件 (*.*)")
    return ";;".join(parts)


def guess_format(path: str) -> str:
    ext = Path(path).suffix.lower().lstrip(".")
    return ext if ext in FORMATS else "txt"
