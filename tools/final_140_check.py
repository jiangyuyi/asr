"""Verify the published 1.4.0: Excel export works from the downloaded build."""
from __future__ import annotations

import os
import shutil
import stat
import subprocess
import tempfile
import time
import zipfile
from pathlib import Path

ROOT = Path(r"D:\Work\asr_mm")
TMP = Path(tempfile.mkdtemp(prefix="asr-mm-140-"))
LOG = ROOT / "tools" / "final_140.txt"
VIDEO = ROOT / "samples/课堂录像示例.mp4"
TAG = "v1.4.0"

lines: list[str] = []
fails: list[str] = []


def say(m=""):
    lines.append(m)
    LOG.write_text("\n".join(lines), encoding="utf-8")


def check(label, ok, detail=""):
    say(f"[{'PASS' if ok else 'FAIL'}] {label}" + (f"  — {detail}" if detail else ""))
    if not ok:
        fails.append(label)


say("=" * 72)
say("RELEASE 1.4.0 VERIFICATION")
say("=" * 72)

for name in (f"asr-mm-{TAG[1:]}-windows-x64.zip",
             f"asr-mm-gui-{TAG[1:]}-macos-arm64.zip"):
    r = subprocess.run(["gh", "release", "download", TAG, "-D", str(TMP),
                        "--clobber", "-p", name], capture_output=True)
    z = TMP / name
    check(f"download {name}", r.returncode == 0 and z.exists(),
          f"{z.stat().st_size / 1_048_576:.0f} MB" if z.exists()
          else r.stderr.decode()[-120:])

# --- macOS bundle keeps its launcher and payload -------------------------
say()
say("--- macOS GUI bundle ---")
with zipfile.ZipFile(TMP / f"asr-mm-gui-{TAG[1:]}-macos-arm64.zip") as zf:
    names = zf.namelist()
    launcher = next((n for n in names if n.endswith(".command")), None)
    check("launcher shipped", launcher is not None, launcher or "")
    if launcher:
        mode = zf.getinfo(launcher).external_attr >> 16
        check("launcher executable", bool(mode & stat.S_IXUSR), oct(mode & 0o777))
    check("engine binaries present",
          any("_internal/runtime/llama-funasr" in n for n in names))
    check("CA bundle present",
          any(n.endswith("cacert.pem") for n in names))
    check("no stray .exe", not any(n.endswith(".exe") for n in names))

# --- the real test: run the published Windows build, write an xlsx --------
say()
say("--- published Windows build, Excel export ---")
with zipfile.ZipFile(TMP / f"asr-mm-{TAG[1:]}-windows-x64.zip") as zf:
    zf.extractall(TMP / "win")
exe = TMP / "win" / "asr-mm" / "asr-mm.exe"
env = dict(os.environ)
env["ASR_MM_HOME"] = str(ROOT / ".asrhome")
env["PYTHONIOENCODING"] = "utf-8"
env["PYTHONUTF8"] = "1"

v = subprocess.run([str(exe), "--version"], capture_output=True, env=env,
                   timeout=180).stdout.decode("utf-8", "replace")
check("version is 1.4.0", "1.4.0" in v, v.strip())

r = subprocess.run(
    [str(exe), "transcribe", str(VIDEO), "-s", "00:00:15", "-e", "00:00:45",
     "-f", "txt,srt,json,xlsx", "-o", str(TMP / "out")],
    capture_output=True, env=env, timeout=900)
out = r.stdout.decode("utf-8", "replace")
check("all four formats written", r.returncode == 0, f"exit {r.returncode}")
books = list((TMP / "out").glob("*.xlsx")) if (TMP / "out").exists() else []
check("workbook produced", bool(books), books[0].name if books else "")

if books:
    from openpyxl import load_workbook
    ws = load_workbook(books[0]).active
    headers = [c.value for c in ws[1]]
    check("headers", headers == ["序号", "开始", "结束", "开始(秒)", "时长(秒)", "内容"],
          str(headers))
    check("one row per segment", ws.max_row == 4, f"{ws.max_row - 1} rows")
    check("absolute timestamps", str(ws.cell(row=2, column=2).value) == "00:00:15.300",
          str(ws.cell(row=2, column=2).value))
    say("")
    say(f"--- {books[0].name} ---")
    for row in ws.iter_rows(values_only=True):
        say("  " + " | ".join("" if x is None else str(x) for x in row))
    say("")

# --- localised workbook headers ------------------------------------------
say("--- workbook headers follow --lang ---")
for lang, first in (("en", "#"), ("ja", "番号")):
    d = TMP / lang
    subprocess.run([str(exe), "--lang", lang, "transcribe", str(VIDEO),
                    "-s", "00:00:15", "-e", "00:00:25", "-f", "xlsx", "-o", str(d)],
                   capture_output=True, env=env, timeout=900)
    got = list(d.glob("*.xlsx")) if d.exists() else []
    if got:
        cell = load_workbook(got[0]).active["A1"].value
        check(f"{lang} header", cell == first, str(cell))
    else:
        check(f"{lang} header", False, "no workbook")

# --- the language warning -------------------------------------------------
r = subprocess.run(
    [str(exe), "--lang", "en", "transcribe", str(VIDEO), "-m", "paraformer",
     "--content-lang", "ja", "-s", "00:00:15", "-e", "00:00:25",
     "--stdout", "--output-format", "srt", "--no-summary"],
    capture_output=True, env=env, timeout=900)
w = r.stderr.decode("utf-8", "replace")
check("model/language mismatch warns", "does not support" in w and "Japanese" in w,
      w.strip()[:100])

# --- GUI still starts -----------------------------------------------------
say()
p = subprocess.Popen([str(TMP / "wingui" / "asr-mm-gui" / "asr-mm-gui.exe"), str(VIDEO)]
                     if (TMP / "wingui").exists() else
                     [str(ROOT / "dist" / "asr-mm-gui" / "asr-mm-gui.exe"), str(VIDEO)],
                     env=env)
alive = 0
for _ in range(12):
    time.sleep(1)
    if p.poll() is not None:
        break
    alive += 1
check("GUI launches", p.poll() is None, f"{alive}s")
if p.poll() is None:
    p.terminate()
    try:
        p.wait(timeout=6)
    except subprocess.TimeoutExpired:
        p.kill()

shutil.rmtree(TMP, ignore_errors=True)
say("")
say(f"RESULT: {len(fails)} failure(s)")
for f in fails:
    say(f"  FAILED: {f}")
print("WROTE", LOG)
