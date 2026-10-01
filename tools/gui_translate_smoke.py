"""GUI 冒烟测试：翻译控件、5 列表格、导出串联是否都接上了。

不开窗口显示，只构造 + 触发 + 断言，跑在无显示环境下。
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from asr_mm import catalog, export  # noqa: E402
from asr_mm.gui.app import (COL_DUR, COL_EN, COL_JA, COL_START,  # noqa: E402
                           COL_TEXT, MainWindow)
from asr_mm.transcribe import Segment, Transcript  # noqa: E402

FAIL: list[str] = []


def check(label: str, ok: bool, detail: str = "") -> None:
    print(f"[{'PASS' if ok else 'FAIL'}] {label}" + (f"  {detail}" if detail else ""))
    if not ok:
        FAIL.append(label)


def main() -> int:
    app = QApplication.instance() or QApplication(sys.argv[:1])
    win = MainWindow()

    # ---- 表格结构
    check("表格 5 列", win.table.columnCount() == 5,
          f"实际 {win.table.columnCount()}")
    labels = [win.table.horizontalHeaderItem(i).text()
              for i in range(win.table.columnCount())]
    check("表头含英文/日文", labels[COL_EN] and labels[COL_JA], " | ".join(labels))

    # ---- 翻译控件
    check("翻译按钮存在", win.btn_translate is not None)
    check("语言复选框 2 个", len(win.chk_tr) == 2, str(list(win.chk_tr)))
    check("取消按钮默认隐藏", not win.btn_tr_cancel.isVisible())
    check("无结果时翻译按钮禁用", not win.btn_translate.isEnabled())

    # ---- 填入假结果，验证列填充
    segs = [Segment(1.0, 2.0, "第一句。"), Segment(2.0, 3.0, "第二句。")]
    tr = Transcript(source="x", source_duration=3.0, range_start=1.0,
                    range_end=3.0, applied_start=1.0, applied_end=3.0,
                    model="nano", model_label="Nano", segments=segs,
                    elapsed=1.0, audio_seconds=2.0)
    win.translations = {"en": ["First.", "Second."], "ja": ["一つ目。", "二つ目。"]}
    win._fill_table(tr)
    check("行数随片段数", win.table.rowCount() == 2)
    check("中文列", win.table.item(0, COL_TEXT).text() == "第一句。")
    check("英文列已填", win.table.item(0, COL_EN).text() == "First.")
    check("日文列已填", win.table.item(0, COL_JA).text() == "一つ目。")
    check("译文列可编辑", win.table.item(0, COL_EN).flags() & Qt.ItemIsEditable)
    check("时间列不可编辑",
          not (win.table.item(0, COL_START).flags() & Qt.ItemIsEditable))

    # ---- 翻译状态刷新
    win.result = tr
    win._refresh_translate_state()
    if catalog.mt_model_ready("en"):
        check("模型就绪时按钮可用", win.btn_translate.isEnabled())
    else:
        check("模型缺失时按钮禁用（带提示）", not win.btn_translate.isEnabled(),
              "tooltip: " + win.btn_translate.toolTip()[:40])

    # ---- 读取表格里的译文（模拟用户手改）
    win.table.item(0, COL_EN).setText("First EDITED.")
    got = win._current_translations()
    check("读到手改后的译文", got["en"][0] == "First EDITED.", str(got.get("en")))
    check("行数与表格一致", len(got["en"]) == 2)

    # ---- 导出串联
    out = Path(__file__).resolve().parent.parent / "build_gui_smoke"
    out.mkdir(exist_ok=True)
    p = out / "smoke.xlsx"
    export.write(p, win._current_segments(), "xlsx", lang="zh",
                 translations=got)
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    rows = list(ws.iter_rows(values_only=True))
    check("导出行数", len(rows) == 3, f"{len(rows)}")
    check("导出 8 列", len(rows[0]) == 8, " | ".join(str(x) for x in rows[0]))
    check("导出译文对齐", rows[1][6] == "First EDITED.", str(rows[1][6]))
    check("导出日文对齐", rows[1][7] == "一つ目。", str(rows[1][7]))

    # ---- 菜单
    check("菜单有下载翻译模型", win.act_download_mt.text() != "")

    print()
    print(f"RESULT: {len(FAIL)} failure(s)" + (f" -> {FAIL}" if FAIL else ""))
    return 1 if FAIL else 0


if __name__ == "__main__":
    raise SystemExit(main())
