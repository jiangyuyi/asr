"""Real run: transcribe, then write every format and read the xlsx back."""
from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ.setdefault("ASR_MM_HOME", r"D:\Work\asr_mm\.asrhome")
ROOT = Path(r"D:\Work\asr_mm")
sys.path.insert(0, str(ROOT))

from asr_mm import export, transcribe  # noqa: E402
from asr_mm.i18n import set_language  # noqa: E402

VIDEO = ROOT / "20250912 哲商現代実験学校 MR1 part1.avi"
OUT = ROOT / "tools" / "export_demo"
LOG = ROOT / "tools" / "export_demo.txt"

lines: list[str] = []
fails: list[str] = []


def say(m=""):
    lines.append(m)
    LOG.write_text("\n".join(lines), encoding="utf-8")


def check(label, ok, detail=""):
    say(f"[{'PASS' if ok else 'FAIL'}] {label}" + (f"  — {detail}" if detail else ""))
    if not ok:
        fails.append(label)


import shutil  # noqa: E402
shutil.rmtree(OUT, ignore_errors=True)

say("=" * 72)
say("EXPORT — real transcript to every format")
say("=" * 72)

tr = transcribe.transcribe(VIDEO, start="00:00:15", end="00:00:45")
say(f"{len(tr.segments)} segments, {tr.elapsed:.2f}s")
say()

written = export.write_all(OUT, "clip", tr.segments, ["txt", "srt", "json", "xlsx"],
                           lang="zh", json_payload=tr.to_dict())
for p in written:
    check(f"wrote {p.name}", p.exists() and p.stat().st_size > 0,
          f"{p.stat().st_size:,} bytes")

# read the workbook back and print it as a table
from openpyxl import load_workbook  # noqa: E402

ws = load_workbook(OUT / "clip.xlsx").active
say()
say(f"--- clip.xlsx  ({ws.max_row} rows x {ws.max_column} cols) ---")
for row in ws.iter_rows(values_only=True):
    say("  " + " | ".join("" if v is None else str(v) for v in row))
say()

headers = [c.value for c in ws[1]]
check("header matches the documented layout",
      headers == ["序号", "开始", "结束", "开始(秒)", "时长(秒)", "内容"], str(headers))
check("one row per segment", ws.max_row == len(tr.segments) + 1,
      f"{ws.max_row - 1} rows vs {len(tr.segments)} segments")
check("timestamps are the video's own, not the slice's",
      ws.cell(row=2, column=2).value == "00:00:15.300",
      str(ws.cell(row=2, column=2).value))
check("text column carries the transcript",
      tr.segments[0].text in str(ws.cell(row=2, column=6).value))

# headers follow the interface language
for lang, first in (("en", "#"), ("ja", "番号")):
    p = OUT / f"clip_{lang}.xlsx"
    export.write(p, tr.segments, "xlsx", lang=lang)
    got = load_workbook(p).active["A1"].value
    check(f"{lang} workbook header", got == first, str(got))

say()
say("--- clip.txt ---")
say((OUT / "clip.txt").read_text(encoding="utf-8")[:300])
say()
say(f"RESULT: {len(fails)} failure(s)")
for f in fails:
    say(f"  FAILED: {f}")
print("WROTE", LOG)
